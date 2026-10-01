"""Bounded local tabular input; header validation happens before pandas mangling."""
import csv
from datetime import date, datetime
from hashlib import sha256
import json
import math
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
import pandas as pd

from app.core.exceptions import InvalidDatasetError


def content_hash(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def object_hash(value: dict[str, Any]) -> str:
    return sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def missing(value: Any) -> bool:
    return value is None or pd.isna(value) or (isinstance(value, str) and not value.strip())


def label_value(value: Any) -> str | None:
    if missing(value):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def feature_value(value: Any, kind: str, location: str) -> Any:
    if isinstance(value, (dict, list, tuple, set)):
        raise InvalidDatasetError(f'{location}: nilai fitur harus skalar, bukan object/list.')
    if missing(value):
        return None
    try:
        if kind == "numeric":
            number = float(value)
            if not math.isfinite(number):
                raise ValueError("not finite")
            return number
        if kind == "date":
            if isinstance(value, (date, datetime)):
                return value.isoformat()[:10]
            text = str(value).strip()
            try:
                return datetime.fromisoformat(text.replace("Z", "+00:00")).date().isoformat()
            except ValueError:
                for pattern in ("%d/%m/%Y", "%d-%m-%Y"):
                    try:
                        return datetime.strptime(text, pattern).date().isoformat()
                    except ValueError:
                        continue
                raise ValueError("date")
        return " ".join(str(value).split())
    except (ValueError, TypeError, OverflowError) as error:
        raise InvalidDatasetError(f"{location}: nilai tidak sesuai tipe {kind}. Tanggal menerima ISO atau DD/MM/YYYY.") from error


def validate_headers(values: list[Any]) -> list[str]:
    headers = [str(value).strip() if value is not None else "" for value in values]
    if not headers or any(not header for header in headers):
        raise InvalidDatasetError("Header kosong; setiap kolom wajib mempunyai nama pada baris pertama.")
    if len(set(headers)) != len(headers):
        raise InvalidDatasetError("Nama header duplikat; gunakan nama kolom yang unik.")
    return headers


def inspect_input(path: Path) -> dict[str, Any]:
    try:
        if path.suffix == ".csv":
            text = path.read_text(encoding="utf-8-sig")
            try:
                delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t").delimiter
                ambiguous = False
            except csv.Error:
                delimiter, ambiguous = ",", True
            return {"sheets": [], "sheet_name": None, "delimiter": delimiter, "delimiter_ambiguous": ambiguous}
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = workbook.sheetnames
        finally:
            workbook.close()
        if not sheets:
            raise InvalidDatasetError("Workbook tidak memiliki sheet.")
        return {"sheets": sheets, "sheet_name": sheets[0], "delimiter": None, "delimiter_ambiguous": False}
    except (OSError, ValueError, UnicodeError, KeyError) as error:
        raise InvalidDatasetError("File rusak atau CSV bukan UTF-8; ekspor ulang sebagai CSV UTF-8.") from error


def read_input(path: Path, sheet_name: str | None = None, delimiter: str | None = None) -> pd.DataFrame:
    try:
        if path.suffix == ".csv":
            separator = delimiter or inspect_input(path)["delimiter"]
            if separator not in {",", ";", "\t"}:
                raise InvalidDatasetError("Pemisah CSV tidak valid.")
            with path.open(encoding="utf-8-sig", newline="") as handle:
                headers = validate_headers(next(csv.reader(handle, delimiter=separator), []))
            frame = pd.read_csv(path, sep=separator, encoding="utf-8-sig", dtype=object, keep_default_na=False, skip_blank_lines=False)
        else:
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                selected = sheet_name or workbook.sheetnames[0]
                if selected not in workbook.sheetnames:
                    raise InvalidDatasetError("Sheet yang dipilih tidak ditemukan.")
                headers = validate_headers(list(next(workbook[selected].values, ())))
            finally:
                workbook.close()
            frame = pd.read_excel(path, sheet_name=selected, dtype=object, keep_default_na=False)
        if len(frame.columns) != len(headers):
            raise InvalidDatasetError("Jumlah kolom tidak sesuai header.")
        frame.columns = headers
        if frame.empty:
            raise InvalidDatasetError("Dataset tidak memiliki baris data.")
        return frame
    except InvalidDatasetError:
        raise
    except Exception as error:
        raise InvalidDatasetError("Dataset tidak dapat dibaca; periksa header, sheet, encoding, dan pemisah CSV.") from error


def suggested_kind(series: pd.Series) -> str:
    values = [value for value in series if not missing(value)]
    if not values:
        return "categorical"
    if all(isinstance(value, (date, datetime)) for value in values):
        return "date"
    try:
        if all(math.isfinite(float(value)) for value in values):
            return "numeric"
    except (ValueError, TypeError, OverflowError):
        pass
    try:
        for value in values:
            feature_value(value, 'date', 'suggestion')
        return 'date'
    except InvalidDatasetError:
        pass
    if sum(len(str(value)) for value in values) / len(values) >= 50:
        return "text"
    return "categorical"
