"""Verify all source values, label coverage, and the upload-to-preprocess path."""
import argparse
from collections import Counter
from hashlib import sha256
import json
from pathlib import Path
import sys
import tempfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from fastapi.testclient import TestClient
from openpyxl import load_workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("labeled", type=Path)
    args = parser.parse_args()
    original = load_workbook(args.source, read_only=True)
    labeled = load_workbook(args.labeled, read_only=True)
    source_rows = list(original["DATA_MENTAH"].values)
    output_rows = list(labeled["DATA_MENTAH"].values)
    assert len(source_rows) == len(output_rows)
    for row_no, (source, output) in enumerate(zip(source_rows, output_rows), 1):
        assert tuple(source[:12]) == tuple(output[:12]), f"Source values differ at Excel row {row_no}"
    counts = dict(Counter(row[12] or "NEEDS_REVIEW" for row in output_rows[1:]))
    original.close()
    labeled.close()
    # The disposable directory belongs only to this verification process.
    import os
    with tempfile.TemporaryDirectory(prefix="maxkomp-prelabeled-check-") as temporary:
        os.environ["STORAGE_ROOT"] = temporary
        from app.core.config import get_settings
        get_settings.cache_clear()
        from app.main import create_app
        with TestClient(create_app()) as client:
            response = client.post("/api/dataset/upload", files={"file": (args.labeled.name, args.labeled.read_bytes())})
            assert response.status_code == 201, response.text
            job_id = response.json()["data"]["job_id"]
            configured = client.put(f'/api/jobs/{job_id}/dataset/configuration', json={'profile': 'maxkomp', 'sheet_name': 'DATA_MENTAH', 'target_column': 'keputusan_laporan', 'features': []})
            assert configured.status_code == 200, configured.text
            overview = client.get(f"/api/jobs/{job_id}").json()["data"]
            assert overview["input_label_source"] == "IMPORTED"
            cleaned = client.post(f"/api/jobs/{job_id}/clean")
            assert cleaned.status_code == 200, cleaned.text
            assert cleaned.json()["data"]["status"] == "LABELED"
            prepared = client.post(f"/api/jobs/{job_id}/preprocess")
            assert prepared.status_code == 200, prepared.text
            stats = client.get(f"/api/jobs/{job_id}/labeling/summary").json()["data"]
            assert stats["unlabeled_records"] == 0
            print(json.dumps({
                "source_rows_preserved": len(source_rows)-1, "input_decisions": counts,
                "cleaning": cleaned.json()["data"], "preprocessing": prepared.json()["data"],
                "cleaned_labels": {key:stats[key] for key in ("selesai_records","tindak_lanjut_records","needs_review_records")},
                "labeled_sha256": sha256(args.labeled.read_bytes()).hexdigest(),
            }))


if __name__ == "__main__":
    main()
